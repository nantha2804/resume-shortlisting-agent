from __future__ import annotations

import os
import re
from io import BytesIO
from typing import List, Dict, Any

from dotenv import load_dotenv

load_dotenv()

try:
    from langchain_groq import ChatGroq
except Exception:  # pragma: no cover
    ChatGroq = None

SKILL_KEYWORDS = [
    "python",
    "sql",
    "aws",
    "docker",
    "kubernetes",
    "langchain",
    "langgraph",
    "fastapi",
    "flask",
    "ai",
    "machine learning",
    "ml",
    "data science",
    "pandas",
    "numpy",
    "tensorflow",
    "pytorch",
    "javascript",
    "typescript",
    "react",
    "node",
    "nodejs",
    "mongodb",
    "postgresql",
    "mysql",
    "redis",
    "spark",
    "airflow",
    "azure",
    "gcp",
    "linux",
    "rest api",
    "api",
    "pytest",
    "ci/cd",
    "git",
    "graphql",
    "microservices",
    "etl",
]


def normalize_text(value: str) -> str:
    if not value:
        return ""
    return re.sub(r"\s+", " ", value.strip().lower())


def extract_resume_text(file_bytes: bytes, filename: str) -> str:
    """Extract text from supported resume formats, including optional image OCR."""
    extension = os.path.splitext(filename)[1].lower()
    if extension == ".txt":
        return file_bytes.decode("utf-8-sig", errors="ignore").strip()
    if extension == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(file_bytes))
            return "\n".join(page.extract_text() or "" for page in reader.pages).strip()
        except Exception as exc:
            raise ValueError(f"Could not read the PDF file: {exc}") from exc
    if extension == ".docx":
        try:
            from docx import Document
            document = Document(BytesIO(file_bytes))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            paragraphs.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            return "\n".join(paragraphs).strip()
        except Exception as exc:
            raise ValueError(f"Could not read the DOCX file: {exc}") from exc
    if extension == ".doc":
        raise ValueError("Legacy .doc files are not supported directly. Save the file as .docx or PDF and try again.")
    if extension in {".png", ".jpg", ".jpeg"}:
        try:
            from PIL import Image
            import pytesseract
            return pytesseract.image_to_string(Image.open(BytesIO(file_bytes))).strip()
        except ImportError as exc:
            raise ValueError("Image OCR requires Pillow, pytesseract, and the Tesseract application.") from exc
        except Exception as exc:
            raise ValueError(f"Could not extract text from the image: {exc}") from exc
    raise ValueError(f"Unsupported file type: {extension or 'unknown'}")


def generate_tailored_resume(resume_text: str, job_description: str) -> str:
    """Generate editable resume content using the candidate's evidence and job skills."""
    matched = extract_skills(resume_text)
    required = extract_skills(job_description)
    relevant = [skill for skill in required if skill in matched]
    return "\n".join([
        "TAILORED RESUME", "", "PROFESSIONAL SUMMARY",
        f"Candidate with experience relevant to {', '.join(relevant[:8]) or 'the target role'}. Resume content has been tailored to emphasize evidence aligned with the job description.", "",
        "CORE SKILLS", ", ".join(sorted(set(matched + relevant))) or "Add verified skills from your experience.", "",
        "RELEVANT EXPERIENCE AND ACHIEVEMENTS", "Review the original resume evidence below and rewrite each bullet with measurable outcomes.",
        resume_text.strip(), "", "TARGETED KEYWORDS", ", ".join(required) or "Add keywords from the job description only when accurate.", "",
        "Note: Verify every statement and add dates, employers, education, and measurable results before applying.",
    ])


def extract_skills(text: str) -> List[str]:
    text_norm = normalize_text(text)
    found = []
    for skill in SKILL_KEYWORDS:
        if skill in text_norm:
            found.append(skill)
    return found


def estimate_score(resume_text: str, job_description: str) -> Dict[str, Any]:
    resume_norm = normalize_text(resume_text)
    jd_norm = normalize_text(job_description)

    required_skills = extract_skills(job_description)
    resume_skills = extract_skills(resume_text)
    matched = sorted(set(required_skills) & set(resume_skills))
    missing = sorted(set(required_skills) - set(resume_skills))

    if not required_skills:
        required_skills = ["python", "sql", "aws", "api"]

    score = round((len(matched) / max(len(required_skills), 1)) * 100)
    score = max(0, min(score, 100))

    if score >= 80:
        decision = "Shortlist"
    elif score >= 55:
        decision = "Human Review"
    else:
        decision = "Reject"

    strengths = []
    for skill in matched[:5]:
        strengths.append(f"Has relevant experience in {skill}.")

    if not strengths:
        strengths = ["Resume does not strongly match the job requirements yet."]

    suggestions = []
    if missing:
        suggestions.append(f"Add more evidence for: {', '.join(missing[:5])}.")
    if "python" not in resume_skills and "python" in required_skills:
        suggestions.append("Mention Python projects or practical implementation examples.")
    if "sql" not in resume_skills and "sql" in required_skills:
        suggestions.append("Add database/query experience with SQL in projects or work history.")
    if "aws" not in resume_skills and "aws" in required_skills:
        suggestions.append("Highlight cloud deployment, hosting, or infrastructure exposure.")
    if not suggestions:
        suggestions.append("Keep experience statements impact-focused and quantify results with numbers.")

    key_points = [
        f"Matched skills: {', '.join(matched[:5]) if matched else 'none'}",
        f"Missing skills: {', '.join(missing[:5]) if missing else 'none'}",
        f"Resume fit: {score}/100",
        f"Recommendation: {decision}",
    ]

    result = {
        "score": score,
        "decision": decision,
        "matched_skills": matched,
        "missing_skills": missing,
        "strengths": strengths,
        "suggestions": suggestions,
        "key_points": key_points,
        "resume_summary": resume_norm[:500],
    }

    return result


def llm_review_resume(resume_text: str, job_description: str) -> Dict[str, Any]:
    if not os.getenv("GROQ_API_KEY") or ChatGroq is None:
        return {} 

    try:
        llm = ChatGroq(model="llama-3.3-70b-versatile", temperature=0.2)
        prompt = (
            "Review this resume against the job description. "
            "Return JSON with keys: score, decision, strengths, suggestions, key_points. "
            "Score should be 0 to 100, decision should be Shortlist/Human Review/Reject.\n\n"
            f"JOB DESCRIPTION:\n{job_description}\n\nRESUME:\n{resume_text}"
        )
        response = llm.invoke(prompt)
        text = getattr(response, "content", str(response))
        return {"llm_feedback": text}
    except Exception:
        return {}


def review_resume(resume_text: str, job_description: str) -> Dict[str, Any]:
    base_result = estimate_score(resume_text, job_description)
    llm_data = llm_review_resume(resume_text, job_description)
    if llm_data:
        base_result.update(llm_data)
    return base_result
